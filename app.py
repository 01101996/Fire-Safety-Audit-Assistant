"""
Fire Safety Audit Assistant
RAG-based (TF-IDF retrieval + optional Claude) audit tool with dashboard and Word report.
"""
import io
import json
import os
import re
from datetime import date
from pathlib import Path

import pandas as pd
import plotly.express as px
import streamlit as st
from docx import Document
from docx.shared import Pt
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

st.set_page_config(page_title="Fire Safety Audit Assistant", page_icon="🔥", layout="wide")

REF_DIR = Path("reference_docs")  # optional: commit reference docs here in GitHub
STATUSES = ["Not assessed", "Compliant", "Non-compliant", "N/A"]
RISKS = ["", "Low", "Medium", "High", "Critical"]
DEFAULT_MODEL = "claude-sonnet-5-5"

CHECKLIST = {
    "Fire detection & alarm": [
        "Fire detection system installed, tested and maintained",
        "Manual call points visible, unobstructed and tested",
        "Alarm audible in all areas; alarm panel free of faults",
    ],
    "Firefighting equipment": [
        "Portable extinguishers correct type, mounted, tagged and in-date",
        "Hose reels / hydrants accessible and serviceable",
        "Sprinkler / suppression system maintained and valves open",
    ],
    "Means of escape": [
        "Exit routes clear, unobstructed and adequately wide",
        "Emergency exits unlocked, open outward, with panic hardware",
        "Exit and directional signage visible and illuminated",
        "Assembly points designated and marked",
    ],
    "Emergency lighting": [
        "Emergency lighting covers escape routes and is function-tested",
    ],
    "Compartmentation": [
        "Fire doors in good condition, self-closing, not wedged open",
        "Penetrations through fire walls/floors fire-stopped",
    ],
    "Electrical & ignition sources": [
        "Electrical panels, cables and sockets free of damage/overloading",
        "Hot work controlled by permit system",
        "Flammable liquids/gases stored and handled safely",
    ],
    "Housekeeping": [
        "Combustible waste and materials controlled",
        "Storage kept clear of ceilings, sprinklers and electrical panels",
    ],
    "Emergency management": [
        "Emergency response plan and fire wardens appointed",
        "Evacuation drills conducted and recorded",
        "Staff fire safety training records available",
        "Fire risk assessment current and reviewed",
        "Equipment inspection and maintenance records available",
    ],
}

COLUMNS = ["ID", "Category", "Checkpoint", "Status", "Risk", "Observation",
           "Reference clause", "AI finding", "Corrective action", "Responsible", "Due date"]


# ----------------------------------------------------------------------------
# Document processing and retrieval
# ----------------------------------------------------------------------------
def extract_text(name: str, data: bytes):
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return [(f"p.{i + 1}", p.extract_text() or "") for i, p in enumerate(reader.pages)]
    if ext == "docx":
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        for t in doc.tables:
            for row in t.rows:
                text += "\n" + " | ".join(c.text for c in row.cells)
        return [("doc", text)]
    return [("txt", data.decode("utf-8", errors="ignore"))]


def chunk_text(source, loc, text, size=200, overlap=40):
    words = text.split()
    out, step = [], size - overlap
    for i in range(0, max(len(words), 1), step):
        piece = " ".join(words[i:i + size])
        if len(piece.split()) >= 15:
            out.append({"source": source, "loc": loc, "text": piece})
    return out


@st.cache_resource(show_spinner="Indexing reference documents...")
def build_index(files: tuple):
    chunks = []
    for name, data in files:
        for loc, text in extract_text(name, data):
            chunks += chunk_text(name, loc, text)
    if not chunks:
        return None
    vec = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
    mat = vec.fit_transform([c["text"] for c in chunks])
    return {"vec": vec, "mat": mat, "chunks": chunks}


def retrieve(query: str, index, k=5):
    if not index:
        return []
    sims = cosine_similarity(index["vec"].transform([query]), index["mat"]).ravel()
    top = sims.argsort()[::-1][:k]
    return [{**index["chunks"][i], "score": float(sims[i])} for i in top if sims[i] > 0]


def format_context(hits):
    return "\n\n".join(f"[{h['source']} {h['loc']}]\n{h['text']}" for h in hits)


# ----------------------------------------------------------------------------
# LLM helpers (optional; app still works in retrieval-only mode)
# ----------------------------------------------------------------------------
def get_api_key():
    try:
        key = st.secrets["ANTHROPIC_API_KEY"]
    except Exception:
        key = None
    return key or st.session_state.get("api_key") or os.getenv("ANTHROPIC_API_KEY")


def call_llm(system: str, user: str, max_tokens=900):
    import anthropic
    client = anthropic.Anthropic(api_key=get_api_key())
    resp = client.messages.create(
        model=st.session_state.get("model", DEFAULT_MODEL),
        max_tokens=max_tokens,
        system=system,
        messages=[{"role": "user", "content": user}],
    )
    return "".join(b.text for b in resp.content if b.type == "text")


ASSESS_SYSTEM = (
    "You are a senior fire safety auditor. Use ONLY the reference excerpts provided. "
    "If they do not support a finding, say so rather than inventing clauses. "
    "Reply with a single JSON object and nothing else, with keys: "
    "clause_reference (document + clause/page), non_compliance_statement (1-2 sentences), "
    "severity (Low|Medium|High|Critical), corrective_action (specific, actionable)."
)


def assess_row(row, index):
    query = f"{row['Category']} {row['Checkpoint']} {row['Observation']}"
    hits = retrieve(query, index, k=5)
    if not hits:
        return None
    user = (f"Checkpoint: {row['Checkpoint']}\nAuditor observation: {row['Observation']}\n\n"
            f"Reference excerpts:\n{format_context(hits)}")
    raw = call_llm(ASSESS_SYSTEM, user)
    match = re.search(r"\{.*\}", raw, re.S)
    data = json.loads(match.group(0)) if match else {}
    return {
        "clause": data.get("clause_reference", ""),
        "finding": data.get("non_compliance_statement", ""),
        "risk": data.get("severity", "") if data.get("severity") in RISKS else "",
        "action": data.get("corrective_action", ""),
    }


# ----------------------------------------------------------------------------
# Audit data
# ----------------------------------------------------------------------------
def new_audit_df():
    rows, n = [], 1
    for cat, items in CHECKLIST.items():
        for item in items:
            rows.append([f"FS-{n:02d}", cat, item, "Not assessed", "", "", "", "", "", "",
                         pd.NaT])
            n += 1
    df = pd.DataFrame(rows, columns=COLUMNS)
    df["Due date"] = pd.to_datetime(df["Due date"])
    return df


def score(df):
    c = (df["Status"] == "Compliant").sum()
    nc = (df["Status"] == "Non-compliant").sum()
    return (100 * c / (c + nc)) if (c + nc) else 0.0, int(c), int(nc)


def rating(pct):
    return "Good" if pct >= 90 else "Satisfactory" if pct >= 75 else "Needs improvement" if pct >= 50 else "Unsatisfactory"


def build_report(df, meta, sources):
    pct, c, nc = score(df)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Fire Safety Audit Report", 0)
    for k, v in meta.items():
        doc.add_paragraph(f"{k}: {v}")
    doc.add_heading("1. Summary", 1)
    doc.add_paragraph(
        f"Checkpoints assessed: {c + nc} | Compliant: {c} | Non-compliant: {nc} | "
        f"Compliance score: {pct:.1f}% ({rating(pct)})")
    doc.add_heading("2. Reference documents", 1)
    for s in sources or ["None loaded"]:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_heading("3. Non-compliances and corrective actions", 1)
    ncdf = df[df["Status"] == "Non-compliant"]
    if ncdf.empty:
        doc.add_paragraph("No non-compliances recorded.")
    else:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Light Grid Accent 1"
        for i, h in enumerate(["ID", "Checkpoint / observation", "Reference", "Risk",
                               "Corrective action", "Owner / due"]):
            table.rows[0].cells[i].text = h
        for _, r in ncdf.iterrows():
            cells = table.add_row().cells
            due = r["Due date"].strftime("%d-%b-%Y") if pd.notna(r["Due date"]) else ""
            cells[0].text = r["ID"]
            cells[1].text = f"{r['Checkpoint']}\n{r['Observation']}\n{r['AI finding']}".strip()
            cells[2].text = r["Reference clause"]
            cells[3].text = r["Risk"]
            cells[4].text = r["Corrective action"]
            cells[5].text = f"{r['Responsible']}\n{due}".strip()
    doc.add_heading("4. Full checklist", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(["ID", "Category", "Checkpoint", "Status"]):
        table.rows[0].cells[i].text = h
    for _, r in df.iterrows():
        cells = table.add_row().cells
        for i, k in enumerate(["ID", "Category", "Checkpoint", "Status"]):
            cells[i].text = str(r[k])
    doc.add_paragraph(
        "\nNote: AI-assisted findings must be verified by a competent auditor against the "
        "applicable standards and local regulations before issue.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# Sidebar
# ----------------------------------------------------------------------------
if "audit" not in st.session_state:
    st.session_state.audit = new_audit_df()
if "chat" not in st.session_state:
    st.session_state.chat = []

with st.sidebar:
    st.title("🔥 Fire Safety Audit")
    st.subheader("Audit details")
    site = st.text_input("Site / facility")
    location = st.text_input("Location")
    auditor = st.text_input("Auditor")
    audit_date = st.date_input("Audit date", date.today())

    st.subheader("Reference documents")
    uploads = st.file_uploader("Upload standards / regulations / SOPs",
                               type=["pdf", "docx", "txt"], accept_multiple_files=True)

    st.subheader("AI settings")
    if get_api_key():
        st.success("API key detected")
    else:
        st.text_input("Anthropic API key (optional)", type="password", key="api_key")
    st.text_input("Model", value=DEFAULT_MODEL, key="model")

files = [(u.name, u.getvalue()) for u in uploads or []]
if REF_DIR.exists():
    files += [(p.name, p.read_bytes()) for p in sorted(REF_DIR.iterdir())
              if p.suffix.lower() in (".pdf", ".docx", ".txt")]
index = build_index(tuple(files)) if files else None
sources = sorted({n for n, _ in files})
meta = {"Site": site, "Location": location, "Auditor": auditor,
        "Audit date": audit_date.strftime("%d-%b-%Y")}

with st.sidebar:
    st.caption(f"{len(sources)} document(s), {len(index['chunks']) if index else 0} passages indexed")
    st.divider()
    saved = st.file_uploader("Resume saved audit (CSV)", type=["csv"], key="resume")
    if saved is not None and st.button("Load saved audit"):
        df = pd.read_csv(saved).fillna("")
        df["Due date"] = pd.to_datetime(df["Due date"].replace("", pd.NaT), errors="coerce")
        st.session_state.audit = df.reindex(columns=COLUMNS)
        st.rerun()

st.title("Fire Safety Audit Assistant")
if not index:
    st.info("Upload reference documents in the sidebar (or place them in `reference_docs/`) "
            "to enable clause-based non-compliance findings.")

tab_audit, tab_ask, tab_dash, tab_report = st.tabs(
    ["📋 Audit checklist", "💬 Ask the standards", "📊 Dashboard", "📄 Report"])

# ----------------------------------------------------------------------------
# Audit checklist
# ----------------------------------------------------------------------------
with tab_audit:
    st.caption("Set the status, add observations, then run AI assessment on non-compliant items.")
    edited = st.data_editor(
        st.session_state.audit,
        hide_index=True,
        use_container_width=True,
        num_rows="dynamic",
        disabled=["ID", "Category", "Checkpoint"],
        column_config={
            "Status": st.column_config.SelectboxColumn(options=STATUSES, required=True),
            "Risk": st.column_config.SelectboxColumn(options=RISKS),
            "Due date": st.column_config.DateColumn(format="DD-MMM-YYYY"),
            "Observation": st.column_config.TextColumn(width="medium"),
            "AI finding": st.column_config.TextColumn(width="large"),
            "Corrective action": st.column_config.TextColumn(width="large"),
        },
        key="editor",
    )
    st.session_state.audit = edited

    c1, c2, c3 = st.columns([1, 1, 2])
    redo = c3.checkbox("Re-assess items that already have an AI finding")
    if c1.button("🤖 AI-assess non-compliances", type="primary"):
        if not index:
            st.error("Upload reference documents first.")
        else:
            df = st.session_state.audit.copy()
            todo = df[(df["Status"] == "Non-compliant") & (redo | (df["AI finding"] == ""))]
            bar = st.progress(0.0)
            for n, (i, row) in enumerate(todo.iterrows(), 1):
                try:
                    if get_api_key():
                        res = assess_row(row, index)
                    else:  # retrieval-only fallback
                        hits = retrieve(f"{row['Category']} {row['Checkpoint']} {row['Observation']}", index, 2)
                        res = {"clause": "; ".join(f"{h['source']} {h['loc']}" for h in hits),
                               "finding": "Relevant clauses retrieved (add API key for AI analysis).",
                               "risk": "", "action": ""} if hits else None
                    if res:
                        df.at[i, "Reference clause"] = res["clause"]
                        df.at[i, "AI finding"] = res["finding"]
                        df.at[i, "Corrective action"] = res["action"] or df.at[i, "Corrective action"]
                        if res["risk"] and not df.at[i, "Risk"]:
                            df.at[i, "Risk"] = res["risk"]
                except Exception as e:
                    st.warning(f"{row['ID']}: {e}")
                bar.progress(n / len(todo))
            st.session_state.audit = df
            st.rerun()
    c2.download_button("💾 Save audit (CSV)", st.session_state.audit.to_csv(index=False),
                       file_name=f"fire_audit_{audit_date}.csv", mime="text/csv")

# ----------------------------------------------------------------------------
# Ask the standards (RAG chat)
# ----------------------------------------------------------------------------
with tab_ask:
    for m in st.session_state.chat:
        with st.chat_message(m["role"]):
            st.markdown(m["content"])
    q = st.chat_input("Ask a question about your reference documents...")
    if q:
        st.session_state.chat.append({"role": "user", "content": q})
        hits = retrieve(q, index, k=5)
        if not hits:
            ans = "No relevant passages found. Upload reference documents or rephrase."
        elif get_api_key():
            try:
                ans = call_llm(
                    "You are a fire safety consultant. Answer ONLY from the excerpts, cite "
                    "[document page] for each point, and say if the excerpts don't cover it.",
                    f"Question: {q}\n\nExcerpts:\n{format_context(hits)}", 1000)
            except Exception as e:
                ans = f"LLM error: {e}"
        else:
            ans = "**Top matching passages** (add an API key for a written answer):\n\n" + "\n\n".join(
                f"*{h['source']} {h['loc']}*: {h['text'][:600]}..." for h in hits[:3])
        st.session_state.chat.append({"role": "assistant", "content": ans})
        st.rerun()

# ----------------------------------------------------------------------------
# Dashboard
# ----------------------------------------------------------------------------
with tab_dash:
    df = st.session_state.audit
    pct, c, nc = score(df)
    m = st.columns(5)
    m[0].metric("Compliance score", f"{pct:.0f}%", rating(pct))
    m[1].metric("Compliant", c)
    m[2].metric("Non-compliant", nc)
    m[3].metric("Not assessed", int((df["Status"] == "Not assessed").sum()))
    m[4].metric("Critical/High risks", int(df["Risk"].isin(["Critical", "High"]).sum()))

    colors = {"Compliant": "#2e9e5b", "Non-compliant": "#d64545", "N/A": "#9aa0a6",
              "Not assessed": "#f0b429"}
    left, right = st.columns(2)
    status_counts = df["Status"].value_counts().reset_index()
    status_counts.columns = ["Status", "Count"]
    left.plotly_chart(px.pie(status_counts, names="Status", values="Count", hole=0.5,
                             color="Status", color_discrete_map=colors,
                             title="Status distribution"), use_container_width=True)
    by_cat = df.groupby(["Category", "Status"]).size().reset_index(name="Count")
    right.plotly_chart(px.bar(by_cat, x="Count", y="Category", color="Status", orientation="h",
                              color_discrete_map=colors, title="Status by category"),
                       use_container_width=True)

    risk_df = df[(df["Status"] == "Non-compliant") & (df["Risk"] != "")]
    if not risk_df.empty:
        rc = risk_df["Risk"].value_counts().reindex(RISKS[1:]).fillna(0).reset_index()
        rc.columns = ["Risk", "Count"]
        st.plotly_chart(px.bar(rc, x="Risk", y="Count", title="Non-compliances by risk",
                               color="Risk", color_discrete_map={
                                   "Low": "#2e9e5b", "Medium": "#f0b429",
                                   "High": "#e8832a", "Critical": "#d64545"}),
                        use_container_width=True)

    open_actions = df[(df["Status"] == "Non-compliant")][
        ["ID", "Checkpoint", "Risk", "Corrective action", "Responsible", "Due date"]].copy()
    st.subheader("Open corrective actions")
    if open_actions.empty:
        st.write("None.")
    else:
        open_actions["Overdue"] = open_actions["Due date"].apply(
            lambda d: "Yes" if pd.notna(d) and d.date() < date.today() else "")
        st.dataframe(open_actions, hide_index=True, use_container_width=True)

# ----------------------------------------------------------------------------
# Report
# ----------------------------------------------------------------------------
with tab_report:
    st.write("Generate the audit report from the current checklist.")
    st.dataframe(st.session_state.audit[st.session_state.audit["Status"] == "Non-compliant"][
        ["ID", "Checkpoint", "Risk", "Reference clause", "Corrective action"]],
        hide_index=True, use_container_width=True)
    st.download_button(
        "📄 Download Word report",
        build_report(st.session_state.audit, meta, sources),
        file_name=f"Fire_Safety_Audit_{audit_date}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary")
    st.caption("AI-assisted findings must be verified by a competent auditor.")
